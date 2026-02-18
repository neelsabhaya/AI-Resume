"""parsers/resume_parser.py

Extracts raw text from PDF and DOCX resume files.
"""
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def parse_pdf(file_path: str) -> str:
    """Extract text from a PDF file.

    Pipeline:
      1. Try standard text extraction via pdfplumber.
      2. If no text is found (common for scanned/image-only PDFs),
         fall back to OCR using pytesseract, if available.
    """
    try:
        import pdfplumber

        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

        # If we got real text, return it
        joined_text = "\n".join(text_parts).strip()
        if joined_text:
            return joined_text

        # Fallback: OCR for scanned/image-only PDFs
        try:
            import pytesseract
            from pytesseract import TesseractNotFoundError

            ocr_parts = []
            with pdfplumber.open(file_path) as pdf_ocr:
                for page in pdf_ocr.pages:
                    try:
                        # Render page to image and run OCR
                        pil_image = page.to_image(resolution=300).original
                        ocr_text = pytesseract.image_to_string(pil_image)
                        if ocr_text and ocr_text.strip():
                            ocr_parts.append(ocr_text)
                    except TesseractNotFoundError:
                        logger.error(
                            "Tesseract OCR engine not found. Install it from "
                            "https://github.com/UB-Mannheim/tesseract/wiki and ensure it is on PATH."
                        )
                        break
                    except Exception as ocr_err:
                        logger.error(f"OCR error on PDF page for {file_path}: {ocr_err}")

            ocr_joined = "\n".join(ocr_parts).strip()
            if ocr_joined:
                logger.info(f"PDF OCR extraction succeeded for {file_path}.")
                return ocr_joined
            else:
                logger.warning(f"OCR found no text for {file_path}.")

        except ImportError:
            logger.error(
                "OCR fallback requires 'pytesseract' (and Tesseract installed). "
                "Install with: pip install pytesseract && install system Tesseract."
            )

        return ""

    except Exception as e:
        logger.error(f"PDF parsing error for {file_path}: {e}")
        return ""


def parse_docx(file_path: str) -> str:
    """Extract text from a DOCX file using python-docx."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
    except Exception as e:
        logger.error(f"DOCX parsing error for {file_path}: {e}")
        return ""


def parse_txt(file_path: str) -> str:
    """Extract text from a plain text file."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as e:
        logger.error(f"TXT parsing error for {file_path}: {e}")
        return ""


def parse_resume(file_path: str) -> str:
    """
    Dispatch to the correct parser based on file extension.

    Args:
        file_path: Absolute or relative path to the resume file.

    Returns:
        Extracted raw text string.
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        text = parse_pdf(str(path))
    elif ext in (".docx", ".doc"):
        text = parse_docx(str(path))
    elif ext == ".txt":
        text = parse_txt(str(path))
    else:
        logger.warning(f"Unsupported file type: {ext}")
        return ""

    if not text.strip():
        logger.warning(f"No text extracted from {file_path}")

    return text.strip()
